"""Parity and defect regressions; every fixture uses disposable staged storage."""
import io
import json
from datetime import datetime

import pdfplumber
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from reportlab.pdfgen.canvas import Canvas

from test_operations import create_material, create_order, catalog
from test_critical_corrections import order_payload


def product(jefe, code='TEST-PROD', **changes):
    payload = dict(codigo=code, cliente='Cliente test', descripcion='Etiquetas test',
                   cant_etiquetas_plancha=10, cant_planchas=100, cant_rollos=1,
                   cant_etiquetas_rollo=1000, largo_plancha_impresa=50,
                   largo_rollo=50, ancho_rollo=0.5, tiempo_base_minutos=1)
    payload.update(changes)
    response = jefe.post('/api/catalogo', json=payload)
    assert response.status_code == 201, response.json
    return response.json['producto']


def order(jefe, **changes):
    payload = dict(cliente='Cliente test', orden_compra='OC-123',
                   fecha_entrega='2026-12-01T12:00',
                   items=[dict(codigo='TEST-PROD', descripcion='Etiquetas test', cantidad=1200)])
    payload.update(changes)
    response = jefe.post('/api/pedidos', json=payload)
    assert response.status_code == 201, response.json
    return response.json['pedido']


def advance(jefe, value, states=('Procesando', 'Terminado', 'Facturado', 'Entregado')):
    for state in states:
        response = jefe.put(f"/api/pedidos/{value['id']}/estado", json={'estado': state})
        assert response.status_code == 200, response.json
        value = response.json['pedido']
    return value


def spreadsheet(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    file = io.BytesIO()
    workbook.save(file)
    file.seek(0)
    return file


def pdf_text(response):
    assert response.status_code == 200, response.get_json(silent=True)
    assert response.mimetype == 'application/pdf'
    with pdfplumber.open(io.BytesIO(response.data)) as document:
        return '\n'.join(page.extract_text() or '' for page in document.pages)


@pytest.mark.parametrize('invalid', [True, {}, 'abc', 1.5, -1])
def test_auto_assignment_rejects_invalid_ids_without_assigning_valid_selection(jefe, invalid):
    value = create_order(jefe)
    value = jefe.get(f"/api/pedidos/{value['id']}").json
    response = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [value['id'], invalid]})
    assert response.status_code == 400, response.json
    assert jefe.get(f"/api/pedidos/{value['id']}").json == value


def test_auto_assignment_rejects_missing_ids_atomically(jefe):
    value = create_order(jefe)
    value = jefe.get(f"/api/pedidos/{value['id']}").json
    response = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [value['id'], 99999]})
    assert response.status_code == 404, response.json
    assert jefe.get(f"/api/pedidos/{value['id']}").json == value


def test_auto_assignment_rejects_nonqueued_selection_atomically(jefe):
    queued, processing = create_order(jefe), create_order(jefe)
    queued = jefe.get(f"/api/pedidos/{queued['id']}").json
    advance(jefe, processing, states=('Procesando',))
    response = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [queued['id'], processing['id']]})
    assert response.status_code == 400, response.json
    assert jefe.get(f"/api/pedidos/{queued['id']}").json == queued


def test_auto_assignment_overflow_rolls_back_entire_batch(jefe):
    product(jefe, tiempo_base_minutos=1e300)
    first, second = create_order(jefe), create_order(jefe)
    first, second = [jefe.get(f"/api/pedidos/{value['id']}").json for value in (first, second)]
    response = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [first['id'], second['id']]})
    assert response.status_code == 400, response.json
    for value in (first, second):
        assert jefe.get(f"/api/pedidos/{value['id']}").json == value


@pytest.mark.parametrize('endpoint', ['auto', 'manual'])
@pytest.mark.parametrize('machines', [('M1', 'M2'), ('Mesp', 'M1')])
def test_assignment_cannot_ignore_incompatible_exclusive_machines(jefe, endpoint, machines):
    product(jefe, code='ONE', maquina_requerida=machines[0])
    product(jefe, code='TWO', maquina_requerida=machines[1])
    value = order(jefe, items=[dict(codigo=code, descripcion=code, cantidad=10) for code in ('ONE', 'TWO')])
    value = jefe.get(f"/api/pedidos/{value['id']}").json
    if endpoint == 'auto':
        response = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [value['id']]})
    else:
        response = jefe.put('/api/produccion/asignar-manual', json={
            'pedido_id': value['id'], 'maquina': machines[0], 'fecha_inicio': '2026-12-01T08:00'})
    assert response.status_code == 400, response.json
    assert jefe.get(f"/api/pedidos/{value['id']}").json == value
    # An invalid order must still be visible in the agenda for correction.
    timeline = jefe.get('/api/produccion/planificador')
    assert timeline.status_code == 200
    assert timeline.json['sin_asignar'][0]['error_planificacion']


@pytest.mark.parametrize('invalid', ['false', 1, None, []])
def test_partial_mode_requires_boolean_without_mutating_order(jefe, invalid):
    value = create_order(jefe)
    value = jefe.get(f"/api/pedidos/{value['id']}").json
    path = f"/api/pedidos/{value['id']}/items/{value['items'][0]['id']}/modo-parcial"
    response = jefe.patch(path, json={'activo': invalid})
    assert response.status_code == 400, response.json
    assert jefe.get(f"/api/pedidos/{value['id']}").json == value


@pytest.mark.parametrize('filters', [{'desde': []}, {'hasta': 2}, {'cliente': {}}, {'estado': True}, {'hasta': '9999-12-31'}])
def test_order_report_pdf_invalid_filters_return_json_error(jefe, filters):
    response = jefe.post('/api/reportes/pedidos/pdf', json=filters)
    assert response.status_code == 400
    assert response.is_json and 'error' in response.json


@pytest.mark.parametrize('url', ['/api/reportes/stock', '/api/reportes/pedidos'])
def test_report_end_date_overflow_is_validation_error(jefe, url):
    response = jefe.get(url + '?hasta=9999-12-31')
    assert response.status_code == 400 and response.is_json


def test_monitor_normalizes_timezone_before_comparing_deadline(jefe, monkeypatch):
    # 13:00 UTC is 08:00 Ecuador: at 09:00 this order is one hour overdue.
    import app.routes.main as routes
    monkeypatch.setattr(routes, 'hora_ecuador', lambda: datetime(2026, 12, 1, 9, 0))
    order(jefe, fecha_entrega='2026-12-01T13:00:00+00:00')
    response = jefe.get('/api/pedidos/monitor-tv')
    assert response.status_code == 200
    assert response.json['en_riesgo'] == 1
    assert response.json['pedidos'][0]['en_riesgo'] is True


def test_delivered_edit_and_delete_restore_actual_history_after_catalog_relink(app, jefe):
    first = create_material(jefe, stock=1)
    catalog(app, first['id'])
    second_response = jefe.post('/api/materiales', json=dict(
        codigo_barras='OTHER-ROLL', nombre='Otro rollo', categoria='Materia Prima', stock_actual=5))
    assert second_response.status_code == 201
    second = second_response.json['material']
    value = advance(jefe, create_order(jefe))
    assert jefe.get('/api/materiales').json[1]['stock_actual'] == 0  # Rollo test
    entry = jefe.get('/api/catalogo').json[0]
    entry['material_id'] = second['id']
    assert jefe.put(f"/api/catalogo/{entry['id']}", json=entry).status_code == 200
    payload = order_payload(value)
    payload['items'][0]['cantidad'] = 2400
    response = jefe.put(f"/api/pedidos/{value['id']}", json=payload)
    assert response.status_code == 200, response.json
    stock = {m['id']: m['stock_actual'] for m in jefe.get('/api/materiales').json}
    assert stock == {first['id']: 1, second['id']: 2}
    assert jefe.delete(f"/api/pedidos/{value['id']}").status_code == 200
    stock = {m['id']: m['stock_actual'] for m in jefe.get('/api/materiales').json}
    assert stock == {first['id']: 1, second['id']: 5}
    movements = jefe.get('/api/movimientos').json
    assert sorted((m['tipo_movimiento'], m['cantidad']) for m in movements) == [
        ('Entrada', 1), ('Entrada', 3), ('Salida', 1), ('Salida', 3)]
    assert jefe.delete(f"/api/pedidos/{value['id']}").status_code == 404
    assert jefe.get('/api/movimientos').json == movements


def test_same_code_resolves_material_for_order_customer_only(app, jefe):
    material = create_material(jefe)
    catalog(app, material['id'])
    product(jefe, cliente='Otro cliente')
    first = create_order(jefe)
    other = order(jefe, cliente='Otro cliente')
    advance(jefe, other)
    assert jefe.get('/api/materiales').json[0]['stock_actual'] == 5
    advance(jefe, first)
    assert jefe.get('/api/materiales').json[0]['stock_actual'] == 3


def test_partial_history_survives_header_edit_toggle_and_delete_reopens_finished_order(jefe):
    value = create_order(jefe, quantity=100)
    path = f"/api/pedidos/{value['id']}"
    item_path = path + f"/items/{value['items'][0]['id']}"
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': True}).status_code == 200
    response = jefe.post(item_path + '/entregas', json={'cantidad_entregada': 40})
    assert response.json['pedido']['progreso'] == 40
    delivery_id = response.json['pedido']['items'][0]['entregas'][0]['id']
    shortage = jefe.post(item_path + '/faltantes', json={'cantidad': 5, 'unidad': 'planchas'}).json['faltante']
    value = jefe.get(path).json
    payload = order_payload(value, fecha_entrega='2026-12-02T08:00')
    response = jefe.put(path, json=payload)
    assert response.status_code == 200
    assert response.json['pedido']['items'][0]['entregas'][0]['id'] == delivery_id
    assert response.json['pedido']['items'][0]['faltantes'][0]['id'] == shortage['id']
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': False}).json['pedido']['progreso'] == 0
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': True}).json['pedido']['progreso'] == 40
    assert jefe.post(item_path + '/entregas', json={'cantidad_entregada': 61}).status_code == 400
    assert jefe.post(item_path + '/entregas', json={'cantidad_entregada': 60}).json['pedido']['progreso'] == 100
    advance(jefe, value, states=('Procesando', 'Terminado', 'Facturado'))
    response = jefe.delete(item_path + f'/entregas/{delivery_id}')
    assert response.status_code == 200
    assert response.json['pedido']['estado'] == 'Procesando'
    assert response.json['pedido']['progreso'] == 60
    assert response.json['pedido']['items'][0]['estado_item'] == 'pendiente'
    assert jefe.put(f"/api/pedidos/faltantes/{shortage['id']}/resolver", json={}).json['pedido']['tiene_faltantes_activos'] is False
    assert jefe.delete(f"/api/pedidos/faltantes/{shortage['id']}").json['pedido']['items'][0]['faltantes'] == []


def test_item_and_delivery_ids_cannot_cross_orders(jefe):
    first, other = create_order(jefe), create_order(jefe)
    other = jefe.get(f"/api/pedidos/{other['id']}").json
    path = f"/api/pedidos/{first['id']}/items/{other['items'][0]['id']}"
    for method, suffix, payload in [('put', '/estado', {'estado': 'listo'}),
                                    ('post', '/entregas', {'cantidad_entregada': 1}),
                                    ('patch', '/modo-parcial', {'activo': True}),
                                    ('post', '/faltantes', {'cantidad': 1})]:
        assert getattr(jefe, method)(path + suffix, json=payload).status_code == 404
    assert jefe.get(f"/api/pedidos/{other['id']}").json == other


def test_duplicate_items_and_no_code_products_remain_distinct_in_mrp(jefe):
    product(jefe, code='-', descripcion='Sin código A')
    product(jefe, code='-', descripcion='Sin código B')
    value = order(jefe, items=[dict(codigo='-', descripcion=desc, cantidad=quantity)
                              for desc, quantity in [('Sin código A', 400), ('Sin código A', 400), ('Sin código B', 200)]])
    result = jefe.post('/api/reportes/mrp-lote', json={'pedidos_ids': [value['id'], value['id']]})
    assert result.status_code == 200, result.json
    assert result.json['totales']['total_pedidos'] == 1
    assert result.json['totales']['total_etiquetas'] == 1000
    assert sorted(i['cantidad_total'] for i in result.json['items_agrupados']) == [200, 800]
    assert len(result.json['materiales_bodega']) == 2
    assert all(m['suficiente_stock'] is None for m in result.json['materiales_bodega'])


def test_manual_assignment_conflicts_boundary_timezone_and_unassignment(jefe):
    first, other = create_order(jefe), create_order(jefe)
    other = jefe.get(f"/api/pedidos/{other['id']}").json
    payload = dict(pedido_id=first['id'], maquina='M1', fecha_inicio='2026-12-01T08:00', duracion_minutos=10)
    assert jefe.put('/api/produccion/asignar-manual', json=payload).status_code == 200
    payload.update(pedido_id=other['id'], fecha_inicio='2026-12-01T13:05:00+00:00')
    assert jefe.put('/api/produccion/asignar-manual', json=payload).status_code == 400
    assert jefe.get(f"/api/pedidos/{other['id']}").json == other
    payload['fecha_inicio'] = '2026-12-01T13:10:00+00:00'
    result = jefe.put('/api/produccion/asignar-manual', json=payload)
    assert result.status_code == 200
    assert result.json['pedido']['fecha_inicio_produccion'] == '2026-12-01T08:10:00'
    assert jefe.put('/api/produccion/asignar-manual', json={'pedido_id': first['id'], 'maquina': None}).status_code == 200
    assert jefe.get('/api/produccion/planificador').json['sin_asignar'][0]['id'] == first['id']


def test_auto_assignment_preserves_unselected_reservations_and_fixed_machines(jefe, monkeypatch):
    import app.routes.main as routes
    monkeypatch.setattr(routes, 'hora_ecuador', lambda: datetime(2026, 12, 1, 8, 0))
    product(jefe, code='FIXED', maquina_requerida='M1')
    reservation = create_order(jefe)
    assert jefe.put('/api/produccion/asignar-manual', json=dict(
        pedido_id=reservation['id'], maquina='M1', fecha_inicio='2026-12-01T08:00', duracion_minutos=10)).status_code == 200
    advance(jefe, reservation, states=('Procesando',))
    reserved = jefe.get(f"/api/pedidos/{reservation['id']}").json
    fixed = order(jefe, items=[dict(codigo='FIXED', descripcion='Fijo', cantidad=10)])
    special = order(jefe, items=[dict(codigo='90006718', descripcion='Especial', cantidad=10)])
    flexible = create_order(jefe, quantity=10)
    result = jefe.post('/api/produccion/auto-asignar', json={'pedidos_ids': [fixed['id'], special['id'], flexible['id']]})
    assert result.status_code == 200, result.json
    assigned = {p['id']: p for p in result.json['pedidos']}
    assert assigned[fixed['id']]['maquina_asignada'] == 'M1'
    assert assigned[fixed['id']]['fecha_inicio_produccion'] == '2026-12-01T08:10:00'
    assert assigned[special['id']]['maquina_asignada'] == 'Mesp'
    assert assigned[flexible['id']]['maquina_asignada'] == 'M2'
    assert assigned[flexible['id']]['minutos_estimados'] == 10
    assert jefe.get(f"/api/pedidos/{reservation['id']}").json == reserved


def test_calculator_and_batch_use_exact_consumption_and_snapshot_pdf(app, jefe):
    material = create_material(jefe)
    catalog(app, material['id'])
    entry = jefe.get('/api/catalogo').json[0]
    payload = {'producto_id': entry['id'], 'cantidad': 1201}
    result = jefe.post('/api/calculadora/planchas', json=payload)
    assert result.status_code == 200, result.json
    value = result.json
    assert value['planchas_totales'] == 121
    assert value['rollos_fraccion'] == '121/100'
    assert value['metros_lineales'] == 60.5
    assert value['metros_cuadrados'] == 30.25
    assert (value['paquetes_200'], value['residuo_empaque'], value['bultos']) == (6, 1, 7)
    assert value['materiales_bodega'][0]['suficiente_stock'] is True
    text = pdf_text(jefe.post('/api/calculadora/planchas/pdf', json={**payload, 'revision_calculo': value['revision_calculo']}))
    assert 'TEST-PROD' in text and '1201' in text
    order_value = create_order(jefe, quantity=1201)
    batch_payload = {'pedidos_ids': [order_value['id']]}
    batch = jefe.post('/api/reportes/mrp-lote', json=batch_payload).json
    assert batch['items_agrupados'][0]['rollos_fraccion'] == '121/100'
    assert 'TEST-PROD' in pdf_text(jefe.post('/api/reportes/mrp-lote/pdf', json={
        **batch_payload, 'revision_calculo': batch['revision_calculo']}))
    # Stock changes are part of both server-reconstructed snapshots.
    jefe.post('/api/movimientos', json={'codigo_barras': 'TEST-ROLLO', 'tipo_movimiento': 'Salida', 'cantidad': 1})
    for url, body, revision in [('/api/calculadora/planchas/pdf', payload, value['revision_calculo']),
                                ('/api/reportes/mrp-lote/pdf', batch_payload, batch['revision_calculo'])]:
        response = jefe.post(url, json={**body, 'revision_calculo': revision})
        assert response.status_code == 409 and response.json['codigo'] == 'calculo_desactualizado'
    assert jefe.get('/api/materiales').json[0]['stock_actual'] == 4
    assert len(jefe.get('/api/movimientos').json) == 1


@pytest.mark.parametrize('quantity', [True, 0, -1, 1.5, 'NaN', 9007199254740992])
def test_calculator_rejects_invalid_quantities_without_stock_impact(jefe, quantity):
    entry = product(jefe)
    result = jefe.post('/api/calculadora/planchas', json={'producto_id': entry['id'], 'cantidad': quantity})
    assert result.status_code == 400
    assert jefe.get('/api/movimientos').json == []


def test_calculator_reports_missing_data_and_unknown_units(jefe):
    incomplete = product(jefe, code='MISSING', cant_etiquetas_plancha=0)
    result = jefe.post('/api/calculadora/planchas', json={'producto_id': incomplete['id'], 'cantidad': 1})
    assert result.status_code == 400 and result.json['datos_faltantes']
    material = jefe.post('/api/materiales', json=dict(codigo_barras='INK', nombre='Tinta', categoria='Tintas', stock_actual=100)).json['material']
    entry = product(jefe, code='INK-PROD', material_id=material['id'])
    result = jefe.post('/api/calculadora/planchas', json={'producto_id': entry['id'], 'cantidad': 1})
    assert result.status_code == 200
    assert result.json['materiales_bodega'][0]['suficiente_stock'] is None
    assert result.json['materiales_bodega'][0]['stock_actual_rollos'] is None


def test_import_is_additive_skips_invalid_quantities_and_export_roundtrips(jefe):
    material = create_material(jefe)
    rows = [['SKU', 'Nombre', 'Categoría', 'Stock Actual', 'Stock Mínimo'],
            ['TEST-ROLLO', 'No reemplaza nombre', 'Tintas', 2, 0],
            ['INK', 'Tinta importada', 'Tintas', 3, 1],
            ['BAD', 'Inválida', 'Materia Prima', -1, 0],
            ['FRACTION', 'Inválida', 'Materia Prima', 1.2, 0],
            [None, 'Sin código', 'Tintas', 3, 0]]
    response = jefe.post('/api/importar/excel', data={'archivo': (spreadsheet(rows), 'inventario.xlsx')})
    assert response.status_code == 200, response.json
    assert (response.json['creados'], response.json['actualizados'], response.json['omitidos']) == (1, 1, 3)
    materials = {m['codigo_barras']: m for m in jefe.get('/api/materiales').json}
    assert materials['TEST-ROLLO']['stock_actual'] == 7
    assert materials['TEST-ROLLO']['nombre'] == material['nombre']
    assert materials['INK']['unidad_medida'] == 'unidades'
    exported = jefe.get('/api/exportar/excel')
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.data), data_only=True)
    values = list(workbook.active.values)
    assert len(values) == 3
    assert {row[0]: row[5] for row in values[1:]} == {'TEST-ROLLO': 7, 'INK': 3}
    invalid = jefe.post('/api/importar/excel', data={'archivo': (spreadsheet([['Otra columna'], ['BAD']]), 'bad.xlsx')})
    assert invalid.status_code == 400
    assert {m['codigo_barras']: m for m in jefe.get('/api/materiales').json} == materials


def test_order_pdf_and_report_filters_include_escaped_notes_and_exclude_other_customer(jefe):
    value = order(jefe, cliente='Cliente <A> & asociados')
    order(jefe, cliente='Otro cliente', orden_compra='OC-EXCLUDE')
    path = f"/api/pedidos/{value['id']}/items/{value['items'][0]['id']}/estado"
    assert jefe.put(path, json={'estado': 'suspendido', 'observacion': 'Falta <tinta> & papel'}).status_code == 200
    text = pdf_text(jefe.get(f"/api/pedidos/{value['id']}/pdf"))
    assert 'OC-123' in text and 'TEST-PROD' in text
    filters = {'cliente': 'Cliente <A>', 'estado': 'En Cola'}
    selected = jefe.get('/api/reportes/pedidos', query_string=filters).json
    assert [p['id'] for p in selected] == [value['id']]
    text = pdf_text(jefe.post('/api/reportes/pedidos/pdf', json=filters))
    assert 'Falta <tinta> & papel' in text
    assert 'OC-EXCLUDE' not in text


def test_document_pdf_extraction_is_readonly_and_rejects_invalid_documents(jefe):
    file = io.BytesIO()
    document = Canvas(file)
    for index, line in enumerate(['Cliente: Audit test', 'OC-456', 'Fecha entrega: 01/12/2026', 'SKU01 Etiquetas test 125']):
        document.drawString(50, 750 - index * 20, line)
    document.save()
    file.seek(0)
    response = jefe.post('/api/pedidos/leer-documento', data={'archivo': (file, 'pedido.pdf')})
    assert response.status_code == 200, response.json
    assert response.json == dict(cliente='Audit test', orden_compra='OC-456', fecha_entrega='2026-12-01',
                                 items=[dict(codigo='SKU01', descripcion='Etiquetas test', cantidad=125)])
    assert jefe.get('/api/pedidos').json == []
    assert jefe.get('/api/clientes').json == []
    for filename in ('bad.pdf', 'bad.exe'):
        assert jefe.post('/api/pedidos/leer-documento', data={'archivo': (io.BytesIO(b'invalid'), filename)}).status_code == 400


def test_all_business_routes_require_session_and_enforce_revocation(app, jefe):
    from app.permisos import MODULOS
    credentials = dict(username='noaccess', password='audit-password', rol='empleado',
                       permisos=dict.fromkeys(MODULOS, 'ninguno'))
    assert jefe.post('/api/usuarios', json=credentials).status_code == 200
    stranger, denied = app.test_client(), app.test_client()
    assert denied.post('/api/login', json=credentials).status_code == 200
    routes = [rule for rule in app.url_map.iter_rules() if rule.endpoint.startswith('main.')]
    assert len(routes) >= 50
    for rule in routes:
        url = str(rule)
        for variable in rule.arguments:
            url = url.replace('<int:' + variable + '>', '99999').replace('<path:' + variable + '>', 'missing.png')
        method = next(method for method in sorted(rule.methods) if method not in ('HEAD', 'OPTIONS'))
        response = stranger.open(url, method=method, json={})
        assert response.status_code == 401, (rule.endpoint, method, response.status_code)
        response = denied.open(url, method=method, json={})
        assert response.status_code == 403, (rule.endpoint, method, response.status_code)


def test_material_crud_duplicate_addition_deletion_guards_and_counters(jefe):
    value = create_material(jefe)
    duplicate = jefe.post('/api/materiales', json=dict(codigo_barras='test-rollo', nombre='Ignored',
                          categoria='Materia Prima', stock_actual=2))
    assert duplicate.status_code == 200
    assert duplicate.json['material']['id'] == value['id']
    assert duplicate.json['material']['stock_actual'] == 7
    response = jefe.put(f"/api/materiales/{value['id']}", json=dict(
        revision=duplicate.json['material']['revision'], nombre='Editado', categoria='Tintas', stock_actual=0))
    assert response.status_code == 200
    assert response.json['material']['unidad_medida'] == 'unidades'
    assert response.json['material']['clasificacion'] == 'Insumo'
    assert jefe.get('/api/materiales/contador-criticos').json == {'contador': 1}
    assert jefe.post('/api/movimientos', json=dict(codigo_barras='TEST-ROLLO', tipo_movimiento='Entrada', cantidad=2)).status_code == 201
    assert jefe.delete(f"/api/materiales/{value['id']}").status_code == 400
    unused = jefe.post('/api/materiales', json=dict(codigo_barras='UNUSED', nombre='Sin uso', categoria='Herramientas')).json['material']
    assert jefe.delete(f"/api/materiales/{unused['id']}").status_code == 200
    assert len(jefe.get('/api/materiales').json) == 1


def test_stock_report_counts_movements_with_inclusive_end_day(jefe, monkeypatch):
    import app.routes.main as routes
    create_material(jefe)
    monkeypatch.setattr(routes, 'hora_ecuador', lambda: datetime(2026, 12, 1, 8, 0))
    # Movimiento's column default is bound at declaration time; use the model
    # clock only for these disposable historical fixtures.
    from app.extensions import db
    from app.models import Movimiento
    for movement, quantity in [('Entrada', 2), ('Salida', 1)]:
        assert jefe.post('/api/movimientos', json=dict(codigo_barras='TEST-ROLLO', tipo_movimiento=movement, cantidad=quantity)).status_code == 201
    with jefe.application.app_context():
        movements = Movimiento.query.order_by(Movimiento.id).all()
        movements[0].fecha = datetime(2026, 12, 1, 23, 59)
        movements[1].fecha = datetime(2026, 12, 2, 0, 0)
        db.session.commit()
    response = jefe.get('/api/reportes/stock?desde=2026-12-01&hasta=2026-12-01')
    assert response.status_code == 200
    assert response.json == [dict(material='Rollo test', sku='TEST-ROLLO', unidad_medida='rollos',
                                 total_entradas=2, total_salidas=0, stock_resultante=6)]


def image_file(width=20, height=20):
    result = io.BytesIO()
    Image.new('RGB', (width, height), 'red').save(result, format='PNG')
    result.seek(0)
    return result


def test_client_crud_logo_validation_and_monitor_reference(jefe):
    payload = dict(codigo='AUDIT-CLIENT', nombre='Cliente logo', ruc='123',
                   direccion_facturacion='Calle A', direccion_entrega='Calle B',
                   contacto_principal='Contacto', horarios_entrega='08:00', plazos_entrega='2 días')
    response = jefe.post('/api/clientes', json=payload)
    assert response.status_code == 201
    value = response.json['cliente']
    path = f"/api/clientes/{value['id']}"
    assert jefe.post('/api/clientes', json=payload).status_code == 400
    assert jefe.put(path, json={**payload, 'contacto_principal': 'Otro contacto'}).status_code == 200
    for data, filename in [(image_file(201, 20), 'large.png'), (io.BytesIO(b'invalid'), 'bad.png'), (image_file(), 'bad.jpg')]:
        assert jefe.post(path + '/logo', data={'archivo': (data, filename)}).status_code == 400
    response = jefe.post(path + '/logo', data={'archivo': (image_file(), 'logo.png')})
    assert response.status_code == 200
    url = response.json['logo_url']
    assert jefe.get(url).mimetype == 'image/png'
    value = order(jefe, cliente=value['id'])
    assert value['cliente'] == payload['nombre']
    assert jefe.get('/api/pedidos/monitor-tv').json['pedidos'][0]['logo_url'] == url
    assert jefe.delete(path).status_code == 200
    assert jefe.get('/api/clientes').json == []
    # Orders store the customer's name; removal must preserve the order.
    assert jefe.get(f"/api/pedidos/{value['id']}").json['cliente'] == payload['nombre']


def test_catalog_code_uniqueness_material_link_and_attachment_lifecycle(jefe, storage):
    material = create_material(jefe)
    value = product(jefe, material_id=material['id'], tipo_codigo='especial')
    duplicate = jefe.post('/api/catalogo', json=dict(codigo='test-prod', cliente=' cliente TEST ', descripcion='Otra'))
    assert duplicate.status_code == 400
    product(jefe, cliente='Otro cliente')
    assert jefe.delete(f"/api/materiales/{material['id']}").status_code == 400
    path = f"/api/catalogo/{value['id']}"
    assert jefe.post(path + '/imagen', data={'archivo': (image_file(), 'bad.exe')}).status_code == 400
    response = jefe.post(path + '/imagen', data={'archivo': (image_file(), 'reference.png')})
    assert response.status_code == 200
    url = response.json['imagen_url']
    assert jefe.get(url).mimetype == 'image/png'
    value = order(jefe)
    assert 'TEST-PROD' in pdf_text(jefe.get(f"/api/pedidos/{value['id']}/pdf"))
    assert jefe.delete(path + '/imagen').status_code == 200
    assert jefe.get(url).status_code == 404
    assert list((storage[1] / 'catalogo').iterdir()) == []
    assert jefe.delete(path).status_code == 200


@pytest.mark.parametrize('production,wanted', [
    ({'cant_etiquetas_plancha': 2, 'cant_etiquetas_rollo': 3, 'cant_planchas': 0, 'cant_rollos': 0}, '4/3'),
    ({'cant_etiquetas_rollo': 0, 'cant_planchas': 9, 'cant_rollos': 2}, '2/9'),
    ({'cant_etiquetas_rollo': 0, 'cant_planchas': 0, 'cant_rollos': 0, 'largo_rollo': 1.25}, '1/2'),
])
def test_consumption_keeps_fractional_ratio_and_geometric_fallbacks(jefe, production, wanted):
    value = product(jefe, **production)
    response = jefe.post('/api/calculadora/planchas', json={'producto_id': value['id'], 'cantidad': 3})
    assert response.status_code == 200, response.json
    assert response.json['rollos_fraccion'] == wanted
    assert response.json['fuente_rendimiento_rollo']


def test_mrp_incompatible_material_dimensions_do_not_assert_stock_availability(jefe):
    material = create_material(jefe)
    product(jefe, code='ONE', material_id=material['id'])
    product(jefe, code='TWO', material_id=material['id'], ancho_rollo=0.75)
    value = order(jefe, items=[dict(codigo=code, descripcion=code, cantidad=100) for code in ('ONE', 'TWO')])
    response = jefe.post('/api/reportes/mrp-lote', json={'pedidos_ids': [value['id']]})
    assert response.status_code == 200
    combined = response.json['materiales_bodega'][0]
    assert combined['rollos_fraccion'] == '1/5'
    assert combined['rollos_necesarios'] == 1
    assert combined['ancho_m'] is None
    assert combined['suficiente_stock'] is None
    assert combined['advertencias']


def test_order_revision_and_catalog_changes_invalidate_snapshot_pdf(jefe):
    entry = product(jefe)
    value = create_order(jefe)
    payload = {'pedidos_ids': [value['id']]}
    initial = jefe.post('/api/reportes/mrp-lote', json=payload).json
    response = jefe.put(f"/api/pedidos/{value['id']}", json=order_payload(value, fecha_entrega='2026-12-02T12:00'))
    assert response.status_code == 200
    assert jefe.post('/api/reportes/mrp-lote/pdf', json={**payload, 'revision_calculo': initial['revision_calculo']}).status_code == 409
    individual = {'producto_id': entry['id'], 'cantidad': 10}
    initial = jefe.post('/api/calculadora/planchas', json=individual).json
    entry['cant_etiquetas_plancha'] = 20
    assert jefe.put(f"/api/catalogo/{entry['id']}", json=entry).status_code == 200
    assert jefe.post('/api/calculadora/planchas/pdf', json={**individual, 'revision_calculo': initial['revision_calculo']}).status_code == 409
    assert jefe.post('/api/calculadora/planchas/pdf', json=individual).status_code == 400


def test_backward_transition_keeps_item_states_and_requires_reason(jefe):
    value = create_order(jefe)
    path = f"/api/pedidos/{value['id']}"
    assert jefe.put(path + '/estado', json={'estado': 'Entregado'}).status_code == 400
    advance(jefe, value, states=('Procesando', 'Terminado', 'Facturado'))
    assert jefe.put(path + '/retroceder', json={}).status_code == 400
    response = jefe.put(path + '/retroceder', json={'observacion': 'Corregir factura'})
    assert response.status_code == 200
    assert response.json['pedido']['estado'] == 'Terminado'
    assert response.json['pedido']['items'][0]['estado_item'] == 'listo'
    assert jefe.get('/api/pedidos/contador-activos').json == {'contador': 1}


def test_manager_self_protection_and_deletion_preserve_movement_history(app, jefe):
    me = jefe.get('/api/sesion').json
    assert jefe.put(f"/api/usuarios/{me['id']}", json={'rol': 'empleado'}).status_code == 400
    assert jefe.delete(f"/api/usuarios/{me['id']}").status_code == 400
    credentials = dict(username='audit-admin', password='audit-password', rol='admin')
    response = jefe.post('/api/usuarios', json=credentials)
    assert response.status_code == 200
    user = response.json['usuario']
    admin = app.test_client()
    assert admin.post('/api/login', json=credentials).status_code == 200
    create_material(jefe)
    assert admin.post('/api/movimientos', json=dict(codigo_barras='TEST-ROLLO', cantidad=1, tipo_movimiento='Salida')).status_code == 201
    value = advance(jefe, create_order(jefe))
    assert admin.put(f"/api/pedidos/{value['id']}", json=order_payload(value)).status_code == 403
    assert jefe.delete(f"/api/usuarios/{user['id']}").status_code == 200
    assert admin.get('/api/sesion').status_code == 401
    assert len(jefe.get('/api/movimientos').json) == 1


@pytest.mark.parametrize('changes', [{'cliente': True}, {'cliente': {'nombre': 'Inválido'}},
                                    {'cliente': ['Inválido']}, {'cliente': 1.5},
                                    {'fecha_entrega': 123}, {'fecha_entrega': {'fecha': '2026-12-01'}}])
@pytest.mark.parametrize('editing', [False, True])
def test_order_rejects_malformed_customer_and_deadline_without_creating_records(jefe, editing, changes):
    original = create_order(jefe) if editing else None
    customers = jefe.get('/api/clientes').json
    payload = order_payload(original) if editing else dict(
        cliente='Cliente inválido', fecha_entrega='2026-12-01T12:00',
        items=[dict(codigo='TEST', descripcion='Test', cantidad=1)])
    payload.update(changes)
    response = (jefe.put(f"/api/pedidos/{original['id']}", json=payload) if editing
                else jefe.post('/api/pedidos', json=payload))
    assert response.status_code == 400, response.json
    assert jefe.get('/api/clientes').json == customers
    assert jefe.get('/api/pedidos').json == ([original] if editing else [])


@pytest.mark.parametrize('kind', ['movement', 'partial', 'assignment'])
def test_concurrent_stock_partial_and_machine_writes_have_one_winner(app, jefe, kind):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    create_material(jefe, stock=1)
    first, second = create_order(jefe, quantity=100), create_order(jefe, quantity=100)
    cookie = jefe.get_cookie('sgp_nueva_session').value
    barrier = Barrier(2)

    def mutate(index):
        worker = app.test_client()
        worker.set_cookie('sgp_nueva_session', cookie)
        barrier.wait(timeout=5)
        if kind == 'movement':
            return worker.post('/api/movimientos', json=dict(
                codigo_barras='TEST-ROLLO', tipo_movimiento='Salida', cantidad=1)).status_code
        if kind == 'partial':
            return worker.post(f"/api/pedidos/{first['id']}/items/{first['items'][0]['id']}/entregas",
                               json={'cantidad_entregada': 60}).status_code
        return worker.put('/api/produccion/asignar-manual', json=dict(
            pedido_id=(first, second)[index]['id'], maquina='M1',
            fecha_inicio='2026-12-01T08:00', duracion_minutos=10)).status_code

    with ThreadPoolExecutor(max_workers=2) as workers:
        statuses = sorted(workers.map(mutate, (0, 1)))
    assert statuses == ([201, 400] if kind == 'movement' else [200, 400])
    if kind == 'movement':
        assert jefe.get('/api/materiales').json[0]['stock_actual'] == 0
        assert len(jefe.get('/api/movimientos').json) == 1
    elif kind == 'partial':
        assert jefe.get(f"/api/pedidos/{first['id']}").json['items'][0]['cantidad_entregada_total'] == 60
    else:
        assert len(jefe.get('/api/produccion/planificador').json['lineas_tiempo']['M1']) == 1


def test_import_invalid_measurement_rolls_back_preceding_rows(jefe):
    rows = [['SKU', 'Nombre', 'Stock Actual', 'Metros por Rollo'],
            ['VALID', 'Válido', 5, 50], ['BAD', 'Inválido', 5, -1]]
    response = jefe.post('/api/importar/excel', data={'archivo': (spreadsheet(rows), 'bad.xlsx')})
    assert response.status_code == 400
    assert jefe.get('/api/materiales').json == []


def test_indicators_positive_counts_match_records_and_manager_identity(app, jefe, monkeypatch):
    from app.extensions import db
    from app.models import (Material, Movimiento, Pedido, PedidoItem, ProductoCatalogo,
                            ReporteFaltante)
    import app.routes.indicators as indicators
    monkeypatch.setattr(indicators, 'hora_ecuador', lambda: datetime(2026, 12, 1, 9))
    me = jefe.get('/api/sesion').json
    with app.app_context():
        materials = [Material(codigo_barras=code, nombre=code, categoria='Materia Prima',
                              stock_actual=stock, estado=state)
                     for code, stock, state in [('CRIT', 1, 'Crítico'), ('EMPTY', 0, 'Agotado'), ('OK', 10, 'Disponible')]]
        db.session.add_all(materials)
        db.session.flush()
        db.session.add_all([Movimiento(material_id=materials[2].id, tipo_movimiento='Entrada', cantidad=1,
                                      fecha=date) for date in (
                                          datetime(2026, 11, 30, 23, 59), datetime(2026, 12, 1),
                                          datetime(2026, 12, 1, 23, 59), datetime(2026, 12, 2))])
        values = [Pedido(cliente='Indicadores', orden_compra=code, estado=state, fecha_entrega=date)
                  for code, state, date in [
                      ('OVERDUE', 'En Cola', '2026-12-01T13:00:00Z'),  # 08:00 local
                      ('PLANNED', 'En Cola', '2026-12-02T08:00'),
                      ('PROCESSING', 'Procesando', '2026-12-02T08:00'),
                      ('DONE', 'Entregado', '2026-01-01T08:00'),
                      ('PICKUP', 'Retiro en Fábrica', '2026-01-01T08:00')]]
        values[1].maquina_asignada = 'M1'
        values[1].fecha_inicio_produccion = datetime(2026, 12, 1, 10)
        item = PedidoItem(codigo='ITEM', descripcion='Dos alertas, un pedido', cantidad=100)
        item.faltantes = [ReporteFaltante(cantidad_faltante=1, resuelto=resolved) for resolved in (False, False, True)]
        values[0].items.append(item)
        db.session.add_all(values)
        complete = dict(cant_etiquetas_plancha=10, cant_planchas=100, cant_rollos=1,
                        largo_plancha_impresa=50, largo_rollo=50, ancho_rollo=0.5)
        db.session.add_all([
            ProductoCatalogo(codigo='COMPLETE', descripcion='Completo', material_id=materials[2].id, **complete),
            ProductoCatalogo(codigo='UNLINKED', descripcion='Sin material', **complete),
            ProductoCatalogo(codigo='MISSING', descripcion='Sin rendimiento', material_id=materials[2].id)])
        db.session.commit()
    response = jefe.get('/api/indicadores')
    assert response.status_code == 200
    assert response.json['cuenta_id'] == me['id']
    assert response.json['permisos'] == me['permisos']
    assert {key: value['cantidad'] for key, value in response.json['contadores'].items()} == dict(
        monitor=4, pedidos=4, produccion=4, gantt=1, preacabados=2, inventario=2,
        operaciones=2, calculadora=3, reportes=4, configuracion=2, alertas=4)
    assert all(isinstance(value['etiqueta'], str) and value['etiqueta'] for value in response.json['contadores'].values())


def test_indicators_permissions_revocation_and_account_switch_do_not_leak_modules(app, jefe):
    from app.permisos import MODULOS
    permissions = dict.fromkeys(MODULOS, 'ninguno')
    permissions.update(inventario='ver')
    response = jefe.post('/api/usuarios', json=dict(username='indicator-reader', password='audit-password',
                         rol='empleado', permisos=permissions))
    assert response.status_code == 200
    reader_id = response.json['usuario']['id']
    reader = app.test_client()
    assert reader.get('/api/indicadores').status_code == 401
    assert reader.post('/api/login', json={'username': 'indicator-reader', 'password': 'audit-password'}).status_code == 200
    create_material(jefe, stock=0)
    create_order(jefe)  # hidden order counters must not be exposed
    response = reader.get('/api/indicadores')
    assert response.status_code == 200
    assert response.json['cuenta_id'] == reader_id
    assert response.json['permisos'] == permissions
    assert {key: value['cantidad'] for key, value in response.json['contadores'].items()} == {'inventario': 1, 'alertas': 1}
    permissions = dict.fromkeys(MODULOS, 'ninguno')
    assert jefe.put(f'/api/usuarios/{reader_id}/permisos', json={'permisos': permissions}).status_code == 200
    response = reader.get('/api/indicadores')
    assert response.json['cuenta_id'] == reader_id
    assert response.json['permisos'] == permissions
    assert set(response.json['contadores']) == {'alertas'}
    assert response.json['contadores']['alertas']['cantidad'] == 0
    assert reader.get('/logout').status_code == 200
    assert reader.get('/api/indicadores').status_code == 401
    assert reader.post('/api/login', json={'username': 'jefe-test', 'password': 'clave-test-12345'}).status_code == 200
    response = reader.get('/api/indicadores')
    assert response.json['cuenta_id'] == jefe.get('/api/sesion').json['id']
    assert response.json['cuenta_id'] != reader_id
    assert response.json['contadores']['pedidos']['cantidad'] == 1


@pytest.mark.parametrize('payload', [[], [{}], '', 'text', 123, True, False, None])
def test_nonobject_json_is_400_on_login_and_writes_but_401_without_session(app, jefe, payload):
    body = json.dumps(payload)
    stranger = app.test_client()
    response = stranger.post('/api/login', data=body, content_type='application/json')
    assert response.status_code == 400 and response.is_json and 'error' in response.json
    for rule in app.url_map.iter_rules():
        if not rule.endpoint.startswith('main.'):
            continue
        methods = rule.methods & {'POST', 'PUT', 'PATCH'}
        if not methods:
            continue
        url = str(rule)
        for argument in rule.arguments:
            url = url.replace('<int:' + argument + '>', '99999').replace('<path:' + argument + '>', 'missing.png')
        for method in methods:
            response = jefe.open(url, method=method, data=body, content_type='application/json')
            assert response.status_code == 400 and response.is_json and 'error' in response.json, (method, url, payload)
            response = stranger.open(url, method=method, data=body, content_type='application/json')
            assert response.status_code == 401 and response.is_json, (method, url, payload)
    assert jefe.get('/api/pedidos').json == []
    assert jefe.get('/api/materiales').json == []
    assert jefe.get('/api/clientes').json == []


@pytest.mark.parametrize('surface', ['created', 'history', 'stock'])
def test_stock_and_movement_rows_include_material_units_without_mixing_rolls_and_inks(jefe, surface):
    create_material(jefe)
    response = jefe.post('/api/materiales', json=dict(codigo_barras='INK', nombre='Tinta', categoria='Tintas', stock_actual=5))
    assert response.status_code == 201
    created = []
    for code in ('TEST-ROLLO', 'INK'):
        response = jefe.post('/api/movimientos', json=dict(codigo_barras=code, tipo_movimiento='Entrada', cantidad=2))
        assert response.status_code == 201
        created.append(response.json['movimiento'])
    rows = created if surface == 'created' else jefe.get('/api/movimientos' if surface == 'history' else '/api/reportes/stock').json
    assert {row['sku']: row.get('unidad_medida') for row in rows} == {'TEST-ROLLO': 'rollos', 'INK': 'unidades'}
    if surface == 'stock':
        assert {row['sku']: (row['total_entradas'], row['stock_resultante']) for row in rows} == {'TEST-ROLLO': (2, 7), 'INK': (2, 7)}
    else:
        assert all(row['cantidad'] == 2 and row['tipo_movimiento'] == 'Entrada' for row in rows)


@pytest.mark.parametrize('endpoint', ['/api/movimientos', '/api/reportes/stock'])
def test_orphaned_movement_units_are_explicitly_unknown(app, jefe, endpoint):
    from app.extensions import db
    from app.models import Movimiento
    with app.app_context():
        db.session.add(Movimiento(material_id=99999, tipo_movimiento='Entrada', cantidad=1))
        db.session.commit()
    response = jefe.get(endpoint)
    assert response.status_code == 200
    assert 'unidad_medida' in response.json[0]
    assert response.json[0]['unidad_medida'] is None


@pytest.mark.parametrize('code', ['SHARED-REF', '-'])
def test_order_reader_gets_customer_safe_detail_references_without_catalog_access(app, jefe, code):
    from app.permisos import MODULOS
    material = create_material(jefe)
    other_material = jefe.post('/api/materiales', json=dict(
        codigo_barras='OTHER-REF', nombre='Otro material', categoria='Materia Prima')).json['material']
    own = product(jefe, code=code, cliente='Cliente test', descripcion='Referencia propia', material_id=material['id'])
    other = product(jefe, code=code, cliente='Otro cliente', descripcion='Referencia ajena', material_id=other_material['id'])
    own_file = jefe.post(f"/api/catalogo/{own['id']}/imagen", data={'archivo': (image_file(), 'own.png')}).json
    other_file = jefe.post(f"/api/catalogo/{other['id']}/imagen", data={'archivo': (image_file(), 'other.png')}).json
    value = order(jefe, cliente='Cliente test', items=[dict(codigo=code, descripcion='Referencia propia', cantidad=10)])
    unrelated = order(jefe, cliente='Sin catálogo', items=[dict(codigo=code, descripcion='Referencia propia', cantidad=10)])
    permissions = dict.fromkeys(MODULOS, 'ninguno')
    permissions['pedidos'] = 'ver'
    response = jefe.post('/api/usuarios', json=dict(username='order-reference-reader', password='audit-password',
                         rol='empleado', permisos=permissions))
    assert response.status_code == 200
    reader = app.test_client()
    assert reader.post('/api/login', json={'username': 'order-reference-reader', 'password': 'audit-password'}).status_code == 200
    assert reader.get('/api/catalogo').status_code == 403
    response = reader.get(f"/api/pedidos/{value['id']}")
    assert response.status_code == 200
    item = response.json['items'][0]
    assert item.get('imagen_url') == own_file['imagen_url']
    assert item.get('imagen_referencia') == own_file['imagen_referencia']
    assert item.get('material_sku') == 'TEST-ROLLO'
    assert item['imagen_url'] != other_file['imagen_url']
    assert reader.get(item['imagen_url']).status_code == 200
    unrelated_item = reader.get(f"/api/pedidos/{unrelated['id']}").json['items'][0]
    assert not unrelated_item.get('imagen_url')
    assert not unrelated_item.get('imagen_referencia')
    assert not unrelated_item.get('material_sku')


def test_detail_references_stop_at_fifty_and_list_monitor_do_not_query_catalog(app, jefe):
    from app.extensions import db
    from sqlalchemy import event
    product(jefe)
    value = order(jefe, items=[dict(codigo='TEST-PROD', descripcion=f'Ítem {index}', cantidad=1) for index in range(55)])
    statements = []

    def capture(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith('SELECT') and 'catalogo_productos' in statement:
            statements.append(statement)

    with app.app_context():
        engine = db.engine
    event.listen(engine, 'before_cursor_execute', capture)
    try:
        response = jefe.get(f"/api/pedidos/{value['id']}")
        assert response.status_code == 200
        assert 0 < len(statements) <= 50
        assert sum('imagen_referencia' in item for item in response.json['items']) == 50
        assert all('imagen_referencia' not in item for item in response.json['items'][50:])
        statements.clear()
        response = jefe.get('/api/pedidos')
        assert response.status_code == 200
        assert all('imagen_referencia' not in item for item in response.json[0]['items'])
        assert jefe.get('/api/pedidos/monitor-tv').status_code == 200
        assert statements == []
    finally:
        event.remove(engine, 'before_cursor_execute', capture)


def test_order_dto_exposes_only_last_stored_backward_reason(jefe):
    value = advance(jefe, create_order(jefe), states=('Procesando', 'Terminado', 'Facturado'))
    path = f"/api/pedidos/{value['id']}"
    assert jefe.put(path + '/retroceder', json={'observacion': 'Revisar factura'}).status_code == 200
    assert jefe.put(path + '/retroceder', json={'observacion': 'Revisar impresión'}).status_code == 200
    response = jefe.get(path)
    assert response.status_code == 200
    assert response.json.get('observacion_retroceso') == 'Revisar impresión'
    assert response.json['estado'] == 'Procesando'


def test_ambiguous_blank_code_catalog_does_not_expose_arbitrary_detail_reference(jefe):
    first = product(jefe, code='-', descripcion='Iguales')
    product(jefe, code='-', descripcion='Iguales')
    assert jefe.post(f"/api/catalogo/{first['id']}/imagen", data={'archivo': (image_file(), 'ambiguous.png')}).status_code == 200
    value = order(jefe, items=[dict(codigo='-', descripcion='Iguales', cantidad=1)])
    response = jefe.get(f"/api/pedidos/{value['id']}")
    assert response.status_code == 200
    assert all(not response.json['items'][0].get(field) for field in ('imagen_url', 'imagen_referencia', 'material_sku'))


def mutate_item_with_expected_state(client, value, action, expected, **preconditions):
    item = value['items'][0]
    path = f"/api/pedidos/{value['id']}/items/{item['id']}"
    payload = {'estado_pedido_esperado': expected, **preconditions}
    if action == 'estado':
        return client.put(path + '/estado', json={**payload, 'estado': 'suspendido', 'observacion': 'Revisión'})
    if action == 'modo':
        return client.patch(path + '/modo-parcial', json={**payload, 'activo': False})
    if action == 'entrega':
        return client.post(path + '/entregas', json={**payload, 'cantidad_entregada': 60})
    return client.delete(path + f"/entregas/{item['entregas'][0]['id']}", json=payload)


@pytest.mark.parametrize('action', ['estado', 'modo', 'entrega', 'eliminar_entrega'])
@pytest.mark.parametrize('finished', [False, True])
def test_expected_order_state_allows_current_processing_and_rejects_finished_stale_editor(jefe, action, finished):
    value = advance(jefe, create_order(jefe, quantity=100), states=('Procesando',))
    path = f"/api/pedidos/{value['id']}"
    item_path = path + f"/items/{value['items'][0]['id']}"
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': True}).status_code == 200
    assert jefe.post(item_path + '/entregas', json={'cantidad_entregada': 20}).status_code == 200
    if finished:
        # A second operator finishes after the first one opened its editor.
        advance(jefe, value, states=('Terminado',))
    before = jefe.get(path).json
    response = mutate_item_with_expected_state(jefe, before, action, 'Procesando')
    if finished:
        assert response.status_code == 409 and response.is_json, response.get_json(silent=True)
        assert response.json['codigo'] == 'estado_pedido_desactualizado'
        assert response.json['estado_actual'] == 'Terminado'
        assert jefe.get(path).json == before
    else:
        assert response.status_code == 200, response.get_json(silent=True)
        after = jefe.get(path).json
        assert after['revision'] > before['revision']
        assert after['estado'] == 'Procesando'
        if action == 'estado':
            assert after['items'][0]['estado_item'] == 'suspendido'
        elif action == 'modo':
            assert after['items'][0]['modo_entregas_parciales'] is False
        else:
            assert after['items'][0]['cantidad_entregada_total'] == (80 if action == 'entrega' else 0)


@pytest.mark.parametrize('expected', [None, 1, True, [], {}, 'No existe', 'procesando'])
def test_expected_order_state_rejects_invalid_provided_value_without_changes(jefe, expected):
    value = advance(jefe, create_order(jefe), states=('Procesando',))
    value = jefe.get(f"/api/pedidos/{value['id']}").json
    response = mutate_item_with_expected_state(jefe, value, 'modo', expected)
    assert response.status_code == 400 and response.is_json, response.get_json(silent=True)
    assert response.json['codigo'] == 'estado_pedido_esperado_invalido'
    assert jefe.get(f"/api/pedidos/{value['id']}").json == value


@pytest.mark.parametrize('mismatch', ['description_case', 'other_customer', 'ambiguous'])
def test_detail_nonmatching_catalog_explicitly_nulls_reference_fields(jefe, mismatch):
    code = 'SHARED-REF' if mismatch == 'other_customer' else '-'
    entry = product(jefe, code=code, descripcion='Referencia exacta', cliente='Cliente test')
    assert jefe.post(f"/api/catalogo/{entry['id']}/imagen", data={'archivo': (image_file(), 'reference.png')}).status_code == 200
    if mismatch == 'ambiguous':
        product(jefe, code=code, descripcion='Referencia exacta', cliente='Cliente test')
    value = order(jefe, cliente='Otro cliente' if mismatch == 'other_customer' else 'Cliente test',
                  items=[dict(codigo=code, descripcion='referencia exacta' if mismatch == 'description_case' else 'Referencia exacta', cantidad=1)])
    response = jefe.get(f"/api/pedidos/{value['id']}")
    assert response.status_code == 200
    item = response.json['items'][0]
    for field in ('imagen_url', 'imagen_referencia', 'material_sku'):
        assert field in item and item[field] is None, (field, item)
        assert field not in jefe.get('/api/pedidos').json[0]['items'][0]


@pytest.mark.parametrize('action', ['estado', 'modo', 'entrega', 'eliminar_entrega'])
@pytest.mark.parametrize('stale', [False, True])
def test_expected_order_revision_guards_all_four_item_mutations(jefe, action, stale):
    value = advance(jefe, create_order(jefe, quantity=100), states=('Procesando',))
    path = f"/api/pedidos/{value['id']}"
    item_path = path + f"/items/{value['items'][0]['id']}"
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': True}).status_code == 200
    assert jefe.post(item_path + '/entregas', json={'cantidad_entregada': 20}).status_code == 200
    captured = jefe.get(path).json
    if stale:
        assert jefe.put(path, json=order_payload(captured, fecha_entrega='2027-01-01T08:00')).status_code == 200
    current = jefe.get(path).json
    response = mutate_item_with_expected_state(jefe, current, action, 'Procesando',
                                              revision_pedido_esperada=captured['revision'])
    if stale:
        assert response.status_code == 409 and response.is_json, response.get_json(silent=True)
        assert response.json['codigo'] == 'revision_desactualizada'
        assert response.json['revision_actual'] == current['revision']
        assert jefe.get(path).json == current
    else:
        assert response.status_code == 200, response.get_json(silent=True)
        assert jefe.get(path).json['revision'] > current['revision']


@pytest.mark.parametrize('intervening', ['suspendido', 'modo_apagado', 'cantidad'])
def test_pending_delivery_cannot_overwrite_new_item_decision_or_quantity(jefe, intervening):
    value = advance(jefe, create_order(jefe, quantity=100), states=('Procesando',))
    path = f"/api/pedidos/{value['id']}"
    item_path = path + f"/items/{value['items'][0]['id']}"
    assert jefe.patch(item_path + '/modo-parcial', json={'activo': True}).status_code == 200
    assert jefe.post(item_path + '/entregas', json={'cantidad_entregada': 20}).status_code == 200
    captured = jefe.get(path).json
    if intervening == 'suspendido':
        response = jefe.put(item_path + '/estado', json={'estado': 'suspendido', 'observacion': 'Revisión nueva'})
    elif intervening == 'modo_apagado':
        response = jefe.patch(item_path + '/modo-parcial', json={'activo': False})
    else:
        payload = order_payload(captured)
        payload['items'][0]['cantidad'] = 200
        response = jefe.put(path, json=payload)
    assert response.status_code == 200
    current = jefe.get(path).json
    assert current['estado'] == 'Procesando'
    assert current['revision'] > captured['revision']
    response = jefe.post(item_path + '/entregas', json=dict(
        cantidad_entregada=80, estado_pedido_esperado='Procesando', revision_pedido_esperada=captured['revision']))
    assert response.status_code == 409, response.get_json(silent=True)
    assert jefe.get(path).json == current
    assert current['items'][0]['cantidad_entregada_total'] == 20


@pytest.mark.parametrize('revision', [None, True, 0, -1, 1.0, '1', {}, []])
def test_expected_order_revision_requires_positive_integer_without_changes(jefe, revision):
    value = advance(jefe, create_order(jefe), states=('Procesando',))
    path = f"/api/pedidos/{value['id']}"
    before = jefe.get(path).json
    response = mutate_item_with_expected_state(jefe, before, 'modo', 'Procesando', revision_pedido_esperada=revision)
    assert response.status_code == 400 and response.is_json, response.get_json(silent=True)
    assert response.json['codigo'] == 'revision_invalida'
    assert jefe.get(path).json == before


def test_revision_precondition_is_independent_of_optional_state_precondition(jefe):
    value = create_order(jefe, quantity=100)
    path = f"/api/pedidos/{value['id']}"
    item_path = path + f"/items/{value['items'][0]['id']}"
    stale_revision = value['revision']
    response = jefe.patch(item_path + '/modo-parcial', json={'activo': True})
    assert response.status_code == 200
    current = jefe.get(path).json
    response = jefe.post(item_path + '/entregas', json={'cantidad_entregada': 10, 'revision_pedido_esperada': stale_revision})
    assert response.status_code == 409, response.get_json(silent=True)
    assert jefe.get(path).json == current
    response = jefe.post(item_path + '/entregas', json={'cantidad_entregada': 10, 'revision_pedido_esperada': current['revision']})
    assert response.status_code == 200
    assert response.json['pedido']['items'][0]['cantidad_entregada_total'] == 10
